from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from moral_dilemmas.analysis import build_final_analysis_workbook


def test_build_final_analysis_workbook_combines_expected_sheets(tmp_path: Path) -> None:
    value_prevalence_path = tmp_path / "value_prevalence_summary.csv"
    bootstrap_prevalence_path = tmp_path / "bootstrap_prevalence_ci.csv"
    contrastive_value_path = tmp_path / "contrastive_value_summary.csv"
    bootstrap_contrastive_path = tmp_path / "bootstrap_contrastive_ci.csv"
    global_glmm_path = tmp_path / "glmm_global_results_adjusted.csv"
    pairwise_path = tmp_path / "glmm_pairwise_results.csv"
    bootstrap_model_differences_path = tmp_path / "bootstrap_model_differences.csv"
    output_path = tmp_path / "final_analysis_results.xlsx"

    pd.DataFrame(
        [
            {
                "model": "model-a",
                "value": "care",
                "available_count": 3,
                "selected_count": 2,
                "prevalence": 0.2,
                "rank": 1,
            },
            {
                "model": "model-b",
                "value": "care",
                "available_count": 3,
                "selected_count": 1,
                "prevalence": 0.1,
                "rank": 1,
            },
        ]
    ).to_csv(value_prevalence_path, index=False)
    pd.DataFrame(
        [
            {
                "model": "model-a",
                "value": "care",
                "prevalence": 0.2,
                "ci_lower": 0.1,
                "ci_upper": 0.3,
            },
            {
                "model": "model-b",
                "value": "care",
                "prevalence": 0.1,
                "ci_lower": 0.0,
                "ci_upper": 0.2,
            },
        ]
    ).to_csv(bootstrap_prevalence_path, index=False)
    pd.DataFrame(
        [
            {
                "model": "model-a",
                "value": "care",
                "contrastive_count": 2,
                "selected_contrastive_count": 1,
                "contrastive_selection_rate": 0.5,
                "eligible_for_glmm": True,
                "contrastive_rank": 1,
            },
            {
                "model": "model-b",
                "value": "care",
                "contrastive_count": 2,
                "selected_contrastive_count": 2,
                "contrastive_selection_rate": 1.0,
                "eligible_for_glmm": True,
                "contrastive_rank": 1,
            },
        ]
    ).to_csv(contrastive_value_path, index=False)
    pd.DataFrame(
        [
            {
                "model": "model-a",
                "value": "care",
                "contrastive_selection_rate": 0.5,
                "ci_lower": 0.2,
                "ci_upper": 0.8,
            },
            {
                "model": "model-b",
                "value": "care",
                "contrastive_selection_rate": 1.0,
                "ci_lower": 0.7,
                "ci_upper": 1.0,
            },
        ]
    ).to_csv(bootstrap_contrastive_path, index=False)
    pd.DataFrame(
        [
            {
                "value": "care",
                "n_dilemmas": 2,
                "n_observations": 8,
                "lr_statistic": 4.0,
                "df": 1,
                "p_value": 0.04,
                "p_adjusted": 0.04,
                "significant": True,
                "convergence_ok": True,
                "singular_fit": False,
            }
        ]
    ).to_csv(global_glmm_path, index=False)
    pd.DataFrame(
        [
            {
                "value": "care",
                "model_1": "model-a",
                "model_2": "model-b",
                "p_value": 0.03,
                "p_adjusted": 0.03,
                "significant": True,
            }
        ]
    ).to_csv(pairwise_path, index=False)
    pd.DataFrame(
        [
            {
                "model_1": "model-a",
                "model_2": "model-b",
                "value": "care",
                "difference": -0.5,
                "ci_lower": -0.8,
                "ci_upper": -0.2,
            }
        ]
    ).to_csv(bootstrap_model_differences_path, index=False)

    summary = build_final_analysis_workbook(
        value_prevalence_path=value_prevalence_path,
        bootstrap_prevalence_path=bootstrap_prevalence_path,
        contrastive_value_path=contrastive_value_path,
        bootstrap_contrastive_path=bootstrap_contrastive_path,
        global_glmm_path=global_glmm_path,
        pairwise_path=pairwise_path,
        bootstrap_model_differences_path=bootstrap_model_differences_path,
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    assert workbook.sheetnames == [
        "value_model_summary",
        "global_glmm_results",
        "pairwise_results",
        "bootstrap_model_differences",
    ]
    assert summary.row_counts == {
        "value_model_summary": 2,
        "global_glmm_results": 1,
        "pairwise_results": 1,
        "bootstrap_model_differences": 1,
    }

    value_sheet = workbook["value_model_summary"]
    headers = [cell.value for cell in value_sheet[1]]
    assert "prevalence_rank" in headers
    assert "prevalence_ci_lower" in headers
    assert "contrastive_ci_upper" in headers
    assert value_sheet.freeze_panes == "A2"
    assert value_sheet.auto_filter.ref is not None
